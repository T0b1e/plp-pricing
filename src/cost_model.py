"""Bottom-up cost of one trip, from the assumptions in 'PLP Cost Structure.xlsx' (same maths as
scripts/build_cost_calculator.py): variable items scale with billable km, fixed items are per trip."""
import re
from functools import lru_cache

import openpyxl

from src.paths import ROOT

COST_WORKBOOK = ROOT / "excel" / "PLP Cost Structure.xlsx"
TRIPS_PER_MONTH = 25   # utilisation that spreads the fixed costs (Assumption row 9)

_ROWS = {"kmpl": 15, "oil_cost": 17, "oil_km": 18, "tyres": 20, "tyre_price": 21, "tyre_km": 23,
         "maint": 25, "genset_lph": 28, "speed": 29, "salary": 35, "allow": 45, "insurance": 37,
         "tax": 38, "other": 46, "overhead": 40, "price": 41, "salvage": 42, "life": 43,
         "threshold": 11, "diesel": 14}


@lru_cache(maxsize=1)
def _assumptions() -> dict[str, dict[str, float]]:
    ws = openpyxl.load_workbook(COST_WORKBOOK, data_only=True)["Assumption"]
    return {ws.cell(5, c).value: {k: ws.cell(r, c).value for k, r in _ROWS.items()} for c in range(4, 14)}


def assumption_class(vehicle_class: str) -> str | None:
    """'6W reefer' -> '6W - Reefer'. None for classes the cost workbook has no column for."""
    m = re.fullmatch(r"(4W|6W|10W) (dry|reefer)", vehicle_class or "")
    return f"{m[1]} - {m[2].capitalize()}" if m else None


def trip_cost(vehicle_class: str, billable_km: float, diesel: float | None = None) -> dict[str, float] | None:
    """THB per trip by cost item at this billable km (round-trip km; one-way = half), or None if unmapped."""
    name = assumption_class(vehicle_class)
    a = _assumptions().get(name) if name else None
    if a is None:
        return None
    diesel = a["diesel"] if diesel is None else diesel
    n = TRIPS_PER_MONTH
    far = 2 if billable_km > a["threshold"] else 1
    genset = a["genset_lph"] * (billable_km / 2 / a["speed"]) * diesel if a["genset_lph"] else 0
    return {
        "Fuel": diesel / a["kmpl"] * billable_km,
        "Oil + filter": a["oil_cost"] / a["oil_km"] * billable_km,
        "Tyres": a["tyres"] * a["tyre_price"] / a["tyre_km"] * billable_km,
        "Maintenance": a["maint"] * billable_km,
        "Driver allowance": a["allow"] * far,
        "Tolls / non-receipt": a["other"] * far,
        "Driver salary": a["salary"] / n,
        "Office overhead": a["overhead"] / n,
        "Insurance": a["insurance"] / (n * 12),
        "Tax + GPS": a["tax"] / (n * 12),
        "Depreciation": a["price"] * (1 - a["salvage"]) / a["life"] / (n * 12),
        "Reefer genset fuel": genset,
    }


def fuel_per_km(vehicle_class: str, diesel: float | None = None) -> tuple[float, float] | None:
    """(diesel THB/L used, running-fuel THB/km) for this class; default diesel is the workbook's."""
    name = assumption_class(vehicle_class)
    a = _assumptions().get(name) if name else None
    if a is None:
        return None
    diesel = a["diesel"] if diesel is None else diesel
    return diesel, diesel / a["kmpl"]
