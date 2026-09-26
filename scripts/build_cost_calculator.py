"""Build 'Cost Reverse Calculator.xlsx' from the assumptions in 'PLP Cost Structure.xlsx'.

Price/km + margin -> cost per trip -> V (THB/km) and F (THB/trip) line items, all as live Excel formulas.
Run: python scripts/build_cost_calculator.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "excel" / "PLP Cost Structure.xlsx"
OUT = ROOT / "excel" / "Cost Reverse Calculator.xlsx"

# Assumption-sheet row -> (label, unit) for the per-class constants we copy.
PARAMS = [
    (15, "Fuel economy", "km/L"),
    (17, "Oil + filter cost per change", "THB"),
    (18, "Oil change interval", "km"),
    (20, "Tyres per vehicle", "count"),
    (21, "Price per tyre", "THB"),
    (23, "Tyre life", "km"),
    (25, "Maintenance", "THB/km"),
    (28, "Reefer genset consumption", "L/hr"),
    (29, "Average speed", "km/hr"),
    (35, "Driver salary", "THB/month"),
    (45, "Driver allowance (base rate)", "THB/trip"),
    (37, "Insurance (cargo + vehicle)", "THB/year"),
    (38, "Annual tax + GPS", "THB/year"),
    (46, "Non-receipt expenses (base rate)", "THB/trip"),
    (40, "Office overhead", "THB/vehicle/month"),
    (41, "Vehicle price", "THB"),
    (42, "Salvage value", "% of price"),
    (43, "Depreciation life", "years"),
    (11, "Allowance threshold (billable km)", "km"),
]
IN_ROW = {src: 4 + i for i, (src, _, _) in enumerate(PARAMS)}  # Assumption row -> Inputs row

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
SECT = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")


def build():
    src = openpyxl.load_workbook(SRC)["Assumption"]
    classes = [src.cell(5, c).value for c in range(4, 14)]
    wb = openpyxl.Workbook()

    # ---- Inputs ---------------------------------------------------------
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = "Per-class assumptions (values copied from PLP Cost Structure.xlsx > Assumption). Edit blue cells."
    inp["A1"].font = BOLD
    inp["A3"], inp["B3"] = "Item", "Unit"
    for j, name in enumerate(classes):
        inp.cell(3, 3 + j, name)
    for c in range(1, 13):
        inp.cell(3, c).font = Font(bold=True, color="FFFFFF")
        inp.cell(3, c).fill = HEAD
    for src_row, label, unit in PARAMS:
        r = IN_ROW[src_row]
        inp.cell(r, 1, f"{label}  (src row {src_row})")
        inp.cell(r, 2, unit)
        for j in range(10):
            c = inp.cell(r, 3 + j, src.cell(src_row, 4 + j).value)
            c.font = BLUE
            c.fill = INPUT
            if src_row == 42:
                c.number_format = "0%"
    inp.column_dimensions["A"].width = 46
    inp.column_dimensions["B"].width = 18
    for j in range(10):
        inp.column_dimensions[col(3 + j)].width = 16
    inp.freeze_panes = "C4"

    def L(src_row):  # per-class lookup, driven by the class index in Calculator!C12
        r = IN_ROW[src_row]
        return f"INDEX(Inputs!$C${r}:$L${r},$C$12)"

    # ---- Calculator -----------------------------------------------------
    ws = wb.create_sheet("Calculator", 0)
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 60
    for c in "EFGH":
        ws.column_dimensions[c].width = 16

    def sect(r, text):
        ws.cell(r, 2, text).font = BOLD
        for c in range(2, 5):
            ws.cell(r, c).fill = SECT

    def row(r, label, formula, note="", fmt="#,##0.00", fill=None):
        ws.cell(r, 2, label)
        c = ws.cell(r, 3, formula)
        c.number_format = fmt
        if fill:
            c.fill = fill
            c.font = BOLD
        ws.cell(r, 4, note).font = Font(italic=True, color="666666")

    ws["B1"] = "Cost reverse calculator: price per km + margin -> cost -> line items"
    ws["B1"].font = Font(bold=True, size=14)
    ws["B2"] = "Margin is on PRICE (price = cost / (1 - margin)), same as the Summary sheet. Yellow cells are inputs."

    sect(3, "1. Inputs")
    inputs = [
        (4, "Vehicle class", classes[0], "@", "pick from list"),
        (5, "One-way distance (km)", 120, "#,##0", ""),
        (6, "Legs billed (2 = round trip)", 2, "0", "Billable km = one-way x legs"),
        (8, "Quoted price per billable km (THB/km)", 20.29, "#,##0.00", "the number you want to reverse-engineer"),
        (9, "Margin on price", 0.30, "0.0%", "Summary sheet uses 30%, Assumption row 49 says 25%"),
        (13, "Does the quoted price already include margin?", "No", "@", "No = quoted is pre-margin, margin gets added on top"),
        (10, "Trips per vehicle per month", 25, "0", "utilisation: spreads fixed cost"),
        (11, "Diesel price (THB/L)", 48, "#,##0.00", ""),
    ]
    for r, label, val, fmt, note in inputs:
        row(r, label, val, note, fmt)
        ws.cell(r, 3).fill = INPUT
        ws.cell(r, 3).font = BLUE
    row(7, "Billable km", "=C5*C6", "", "#,##0")
    row(12, "(class index)", f"=MATCH(C4,Inputs!$C$3:$L$3,0)", "helper", "0")
    dv = DataValidation(type="list", formula1="=Inputs!$C$3:$L$3", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C4")
    dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("C13")

    sect(14, "2. Price -> implied cost (exact)")
    row(15, "Quoted amount per trip", "=C8*C7", "quoted/km x billable km", "#,##0.00")
    row(16, "Margin inside the quote (THB)", '=IF(C13="Yes",C15*C9,0)', "0 when the quote is pre-margin", "#,##0.00")
    row(17, "Implied cost per trip", "=C15-C16", "Yes: quote x (1 - margin).  No: the quote itself", "#,##0.00", KEY)
    row(18, "Implied cost per km", "=C17/C7", "", "#,##0.00", KEY)
    row(19, "Selling price per km incl. margin", '=IF(C13="Yes",C8,C18/(1-C9))', "No: implied cost / (1 - margin)", "#,##0.00", KEY)

    sect(20, "3. V: variable cost, THB per km (scales with distance)")
    row(21, "Fuel", f"=C11/{L(15)}", "diesel price / km per litre", "#,##0.000")
    row(22, "Oil + filter", f"={L(17)}/{L(18)}", "cost per change / change interval", "#,##0.000")
    row(23, "Tyres", f"={L(20)}*{L(21)}/{L(23)}", "tyres x price / tyre life", "#,##0.000")
    row(24, "Maintenance", f"={L(25)}", "assumption from workbook", "#,##0.000")
    row(25, "V  (total variable, THB/km)", "=SUM(C21:C24)", "the slope", "#,##0.000", KEY)

    sect(27, "4. F: fixed cost per trip, THB (does not scale with km)")
    row(28, "Driver allowance", f"={L(45)}*IF(C7>{L(11)},2,1)", "base rate, doubled if billable km > threshold")
    row(29, "Non-receipt expenses (tolls etc.)", f"={L(46)}*IF(C7>{L(11)},2,1)", "base rate, doubled if billable km > threshold")
    row(30, "Driver salary", f"={L(35)}/C10", "per month / trips per month")
    row(31, "Office overhead", f"={L(40)}/C10", "per month / trips per month")
    row(32, "Insurance", f"={L(37)}/(C10*12)", "per year / trips per year")
    row(33, "Annual tax + GPS", f"={L(38)}/(C10*12)", "per year / trips per year")
    row(34, "Depreciation", f"={L(41)}*(1-{L(42)})/{L(43)}/(C10*12)", "(price - salvage) / life / trips per year")
    row(35, "Reefer genset fuel", f"=IF({L(28)}>0,{L(28)}*(C5/{L(29)})*C11,0)", "L/hr x hours (one-way km / speed) x diesel price")
    row(36, "F  (total fixed, THB/trip)", "=SUM(C28:C35)", "the intercept", "#,##0.00", KEY)

    sect(38, "5. Model cost vs implied cost:  cost = V x km + F")
    row(39, "Model cost per trip", "=C25*C7+C36", "V x km + F", "#,##0.00", KEY)
    row(40, "Model cost per km", "=C39/C7", "", "#,##0.00")
    row(41, "Gap: implied cost - model cost (THB/trip)", "=C17-C39", "> 0: price implies MORE cost than the model explains", "#,##0.00")
    row(42, "Margin actually earned if model cost is true", "=1-C39/(C19*C7)", "1 - model cost / selling price", "0.0%", KEY)
    row(43, "Price/km needed for the target margin", "=C39/(1-C9)/C7", "model cost / (1 - margin) / km", "#,##0.00")

    sect(45, "6. What would have to be true to close the gap (one lever at a time)")
    other = "(C39-C11*(C7/{kmpl}+IF({lph}>0,{lph}*C5/{spd},0)))".format(kmpl=L(15), lph=L(28), spd=L(29))
    coeff = "(C7/{kmpl}+IF({lph}>0,{lph}*C5/{spd},0))".format(kmpl=L(15), lph=L(28), spd=L(29))
    row(46, "Diesel price (THB/L) that reproduces implied cost", f"=(C17-{other})/{coeff}", "solves fuel (+ genset) for the gap", "#,##0.00")
    fixed_m = f"({L(35)}+{L(40)}+({L(37)}+{L(38)}+{L(41)}*(1-{L(42)})/{L(43)})/12)"
    a = "(C28+C29+C35)"
    row(47, "Trips/month that reproduces implied cost",
        f'=IF(C17-C25*C7-{a}>0,{fixed_m}/(C17-C25*C7-{a}),"n/a: implied cost below variable+trip costs")',
        "salary, overhead, insurance, tax, depreciation depend on utilisation", "#,##0.0")
    row(48, "Share of model cost that is fixed (F / cost)", "=C36/C39", "the higher, the more utilisation matters", "0.0%")

    sect(50, "7. Cost curve for this class (V x km + F at each distance)")
    heads = ["Billable km", "Cost / trip", "Cost / km", "Price / trip", "Price / km"]
    for j, h in enumerate(heads):
        c = ws.cell(51, 2 + j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center")
    for i, km in enumerate(range(50, 651, 50)):
        r = 52 + i
        ws.cell(r, 2, km).number_format = "#,##0"
        ws.cell(r, 2).alignment = Alignment(horizontal="left")
        ws.cell(r, 3, (
            f"=$C$25*B{r}"
            f"+{L(45)}*IF(B{r}>{L(11)},2,1)+{L(46)}*IF(B{r}>{L(11)},2,1)"
            f"+SUM($C$30:$C$34)"
            f"+IF({L(28)}>0,{L(28)}*(B{r}/$C$6/{L(29)})*$C$11,0)"
        )).number_format = "#,##0"
        ws.cell(r, 4, f"=C{r}/B{r}").number_format = "#,##0.00"
        ws.cell(r, 5, f"=C{r}/(1-$C$9)").number_format = "#,##0"
        ws.cell(r, 6, f"=E{r}/B{r}").number_format = "#,##0.00"
    ws.freeze_panes = "A3"
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
